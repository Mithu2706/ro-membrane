"""API, persistence, authorization, export and report tests."""
import io

from openpyxl import load_workbook

from app import db as dbmod
from app.config import settings

from .conftest import BWRO, SWRO


def test_health_and_catalog(client):
    assert client.get("/api/health").json()["status"] == "ok"
    s = client.get("/api/catalog/summary").json()
    assert s["membranes"] >= 15
    assert s["verified_datasheets"] >= 5


def test_assessment_roundtrip_and_report(client):
    r = client.post("/api/assessments", json={**SWRO, "client": {"name": "Test Client", "project_name": "Pilot"}})
    assert r.status_code == 201, r.text
    body = r.json()
    token = body["reference"]
    assert body["result"]["outcome"] == "recommended"
    assert "technical" not in body["result"]
    g = client.get(f"/api/assessments/{token}")
    assert g.status_code == 200
    assert g.json()["result"]["headline"] == body["result"]["headline"]
    pdf = client.get(f"/api/assessments/{token}/report.pdf")
    assert pdf.status_code == 200
    assert pdf.headers["content-type"] == "application/pdf"
    assert pdf.content[:5] == b"%PDF-"
    assert len(pdf.content) > 5000


def test_invalid_input_returns_422(client):
    assert client.post("/api/assessments", json={**BWRO, "recovery_pct": 0}).status_code == 422
    assert client.post("/api/assessments", json={**BWRO, "feed_tds_mg_l": -5}).status_code == 422
    assert client.post("/api/assessments", json={**BWRO, "application": "nonsense"}).status_code == 422


def test_unknown_token_404(client):
    assert client.get("/api/assessments/does-not-exist").status_code == 404


def test_admin_requires_auth(client):
    for path in ("/api/admin/stats", "/api/admin/submissions", "/api/admin/membranes", "/api/admin/export/submissions.csv",
                 "/api/admin/export/workbook.xlsx", "/api/admin/audit"):
        assert client.get(path).status_code == 401, path
    assert client.post("/api/admin/backup").status_code == 401
    assert client.post("/api/admin/membranes", json={}).status_code == 401


def test_admin_bad_password(client):
    r = client.post("/api/admin/login", json={"username": "admin", "password": "wrong"})
    assert r.status_code == 401


def test_csrf_required_for_state_changes(admin_client):
    token = admin_client.headers.pop("X-CSRF-Token")
    r = admin_client.post("/api/admin/sources", json={"title": "CSRF test", "source_type": "other"})
    assert r.status_code == 403
    admin_client.headers["X-CSRF-Token"] = token
    r = admin_client.post("/api/admin/sources", json={"title": "CSRF test", "source_type": "other"})
    assert r.status_code == 201


def test_admin_submissions_and_detail(admin_client):
    admin_client.post("/api/assessments", json={**BWRO, "client": {"organization": "SearchableOrg"}})
    lst = admin_client.get("/api/admin/submissions", params={"q": "SearchableOrg"}).json()
    assert lst["total"] >= 1
    sid = lst["items"][0]["id"]
    d = admin_client.get(f"/api/admin/submissions/{sid}").json()
    assert "technical" in d["result"]  # admin sees calculation details
    assert admin_client.get(f"/api/admin/submissions/{sid}/report.pdf").content[:5] == b"%PDF-"


def test_membrane_crud_spec_validation_and_audit(admin_client):
    r = admin_client.post("/api/admin/membranes", json={"manufacturer": "TestCo", "model": "CRUD-1", "category": "BWRO",
                                                        "intended_applications": ["bwro"]})
    assert r.status_code == 201, r.text
    mid = r.json()["id"]
    assert "Maximum operating pressure" in r.json()["missing_critical"]
    # wrong unit rejected
    bad = admin_client.post(f"/api/admin/membranes/{mid}/specifications",
                            json={"parameter": "max_operating_pressure_bar", "value": 600, "unit": "psi",
                                  "evidence_type": "manufacturer-rated"})
    assert bad.status_code == 422
    ok = admin_client.post(f"/api/admin/membranes/{mid}/specifications",
                           json={"parameter": "max_operating_pressure_bar", "value": 41, "evidence_type": "manufacturer-rated",
                                 "value_text": "600 psig (41 bar)"})
    assert ok.status_code == 201
    spec_id = ok.json()["id"]
    v = admin_client.post(f"/api/admin/specifications/{spec_id}/verify", json={"verification_status": "verified"})
    assert v.json()["verification_status"] == "verified"
    a = admin_client.post(f"/api/admin/membranes/{mid}/archive", json={"archived": True})
    assert a.json()["archived"] is True
    audit = admin_client.get("/api/admin/audit").json()
    assert any(e["entity"] == "specification" and e["action"] == "verify" for e in audit)


def test_csv_import(admin_client):
    csv_text = ("manufacturer,model,parameter,value,unit,evidence_type,verification_status\n"
                "Hydranautics (Nitto),SWC4+,test_ph,7,pH,manufacturer-rated,unverified\n"
                "Nobody,NOPE,test_ph,7,pH,manufacturer-rated,unverified\n")
    r = admin_client.post("/api/admin/import/specifications", params={"dry_run": "true"},
                          files={"file": ("specs.csv", csv_text, "text/csv")})
    body = r.json()
    assert body["valid_rows"] == 1 and len(body["errors"]) == 1 and body["imported"] == 0


def test_csv_and_excel_export(admin_client):
    admin_client.post("/api/assessments", json=BWRO)
    r = admin_client.get("/api/admin/export/submissions.csv")
    assert r.status_code == 200
    text = r.content.decode("utf-8-sig")
    assert text.splitlines()[0].startswith("id,created_at")
    assert len(text.splitlines()) >= 2
    r = admin_client.get("/api/admin/export/specifications.csv")
    assert "max_operating_pressure_bar" in r.content.decode("utf-8-sig")
    x = admin_client.get("/api/admin/export/workbook.xlsx")
    wb = load_workbook(io.BytesIO(x.content))
    assert {"Submissions", "Membranes", "Specifications", "Observations", "Sources"} <= set(wb.sheetnames)


def test_backup(admin_client):
    r = admin_client.post("/api/admin/backup")
    if settings.database_url:  # hosted Postgres: backups belong to the provider
        assert r.status_code == 501
        return
    assert r.status_code == 200
    assert r.content[:16] == b"SQLite format 3\x00"
    assert any(settings.backup_dir.glob("ro_membrane_*.db"))


def test_database_persists_across_restart(client, tmp_dir):
    token = client.post("/api/assessments", json=BWRO).json()["reference"]
    # Simulate an application restart: dispose the engine and reconnect to the same database
    dbmod.reconnect()
    from fastapi.testclient import TestClient

    from app.main import app

    with TestClient(app) as c2:
        assert c2.get(f"/api/assessments/{token}").status_code == 200
    from sqlalchemy import text

    with dbmod.engine.connect() as con:
        n = con.execute(text("select count(*) from client_submissions")).scalar()
        v = con.execute(text("select max(version) from schema_version")).scalar()
    assert n >= 1 and v == 3


def _pdf_text(content: bytes) -> str:
    import pypdf

    return " ".join(p.extract_text() for p in pypdf.PdfReader(io.BytesIO(content)).pages)


def test_pdf_and_dashboard_are_consistent(client):
    body = client.post("/api/assessments", json={**BWRO, "input_origin": "demo"}).json()
    r = body["result"]
    text = " ".join(_pdf_text(client.get(f"/api/assessments/{body['reference']}/report.pdf").content).split())
    rec = r["recommended"]
    assert r["headline"] in text
    assert rec["status_label"] in text
    e = rec["estimates"]
    assert f"{e['required_pressure_bar']:.1f}" in text and f"{e['required_pressure_high_bar']:.1f}" in text
    for c in r["candidates"]:
        assert c["status_label"] in text
    # Demo origin is visible in both outputs
    assert r["input_origin"] == "demo"
    assert r["warnings"][0]["title"] == "Demonstration data"
    assert "Demonstration data" in text and "Demonstration example" in text
    # Every client warning title appears in the report
    for w in r["warnings"]:
        assert w["title"] in text


def test_report_renders_without_system_fonts(client, monkeypatch):
    """Serverless Linux hosts have no Arial/DejaVu: the report must fall back to ReportLab's bundled Vera fonts."""
    from app.services import report

    monkeypatch.setattr(report, "FONTS", {})
    monkeypatch.setattr(report, "UNICODE_OK", True)
    monkeypatch.setattr(report, "VERA_IN_USE", False)
    monkeypatch.setattr(report, "_FONT_CANDIDATES", {k: [v[-1]] for k, v in report._FONT_CANDIDATES.items()})
    token = client.post("/api/assessments", json=SWRO).json()["reference"]
    pdf = client.get(f"/api/assessments/{token}/report.pdf")
    assert pdf.status_code == 200 and pdf.content[:5] == b"%PDF-"
    assert report.VERA_IN_USE and report.UNICODE_OK
    text = _pdf_text(pdf.content)
    assert "π" in text and "m³" in text  # Unicode preserved with the bundled font
    monkeypatch.setattr(report, "FONTS", {})  # let later tests re-register the platform fonts
