"""CSV and Excel export of research records and client submissions."""
from __future__ import annotations

import csv
import io
import json
from collections.abc import Iterable

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import ClientSubmission, ExperimentalObservation, Membrane, MembraneSpecification, ResearchSource

DATASETS = ("submissions", "membranes", "specifications", "observations", "sources")


def _cell(v):
    if isinstance(v, (dict, list)):
        return json.dumps(v, ensure_ascii=False)
    if hasattr(v, "isoformat"):
        return v.isoformat()
    return v


def rows_for(db: Session, dataset: str) -> tuple[list[str], list[list]]:
    if dataset == "submissions":
        cols = ["id", "created_at", "client_name", "client_organization", "client_email", "project_name", "application",
                "outcome", "recommended_membrane", "feed_tds_mg_l", "feed_temperature_c", "available_pressure_bar",
                "permeate_flow_m3d", "recovery_pct", "target_permeate_tds_mg_l", "warnings", "engine_version"]
        out = []
        for s in db.scalars(select(ClientSubmission).order_by(ClientSubmission.created_at.desc())):
            i = s.inputs or {}
            out.append([s.id, s.created_at, s.client_name, s.client_organization, s.client_email, s.project_name,
                        s.application, s.outcome, s.recommended_membrane, i.get("feed_tds_mg_l"),
                        i.get("feed_temperature_c"), i.get("available_pressure_bar"), i.get("permeate_flow_m3d"),
                        i.get("recovery_pct"), i.get("target_permeate_tds_mg_l"),
                        "; ".join(w.get("title", "") for w in (s.warnings or [])), s.engine_version])
        return cols, out
    if dataset == "membranes":
        cols = ["id", "manufacturer", "model", "category", "intended_applications", "element_size", "feed_limitations",
                "verification_status", "archived", "notes"]
        return cols, [[getattr(m, c) for c in cols] for m in db.scalars(select(Membrane).order_by(Membrane.id))]
    if dataset == "specifications":
        cols = ["id", "membrane_id", "manufacturer", "model", "parameter", "value", "value_min", "value_max", "value_text",
                "unit", "evidence_type", "verification_status", "test_conditions", "source_id", "source_title", "notes"]
        out = []
        for s in db.scalars(select(MembraneSpecification).order_by(MembraneSpecification.membrane_id, MembraneSpecification.id)):
            out.append([s.id, s.membrane_id, s.membrane.manufacturer, s.membrane.model, s.parameter, s.value, s.value_min,
                        s.value_max, s.value_text, s.unit, s.evidence_type, s.verification_status, s.test_conditions,
                        s.source_id, s.source.title if s.source else None, s.notes])
        return cols, out
    if dataset == "observations":
        cols = [c.name for c in ExperimentalObservation.__table__.columns]
        return cols, [[getattr(o, c) for c in cols] for o in db.scalars(select(ExperimentalObservation).order_by(ExperimentalObservation.id))]
    if dataset == "sources":
        cols = ["id", "title", "authors", "year", "doi", "url", "source_type", "publisher", "page_table",
                "verification_status", "notes"]
        return cols, [[getattr(s, c) for c in cols] for s in db.scalars(select(ResearchSource).order_by(ResearchSource.id))]
    raise ValueError(f"Unknown dataset {dataset}")


def to_csv(cols: list[str], rows: Iterable[list]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(cols)
    for r in rows:
        w.writerow([_cell(v) for v in r])
    return buf.getvalue().encode("utf-8-sig")


def to_xlsx(db: Session, datasets: Iterable[str]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill("solid", fgColor="102A2A")
    for ds in datasets:
        cols, rows = rows_for(db, ds)
        ws = wb.create_sheet(ds.capitalize())
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True, color="F5F2EA")
            c.fill = header_fill
        for r in rows:
            ws.append([_cell(v) if not hasattr(v, "tzinfo") else v.replace(tzinfo=None) for v in r])
        for i, col in enumerate(cols, start=1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = min(48, max(10, len(col) + 4))
        ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
